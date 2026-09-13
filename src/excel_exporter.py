"""Excel workbook writer."""
from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADERS = ["Date Discovered", "Time Discovered", "Posted Date", "Job Title", "Location", "Category", "Employment Type", "Req ID", "Tesla Job URL", "Apply?"]


def _date_only(value: object) -> str:
    """Display legacy timestamps and new discovery dates consistently."""
    return str(value)[:10]


def _time_only(value: object) -> str:
    """Extract the local time of day from a discovery timestamp, if present.

    Baseline and legacy rows store a date only, so they have no recorded time
    and yield an empty cell.
    """
    text = str(value)
    if "T" not in text:
        return ""
    time_part = text.split("T", 1)[1]
    # Drop any timezone suffix (e.g. "Z", "+05:30", "-07:00").
    for separator in ("Z", "+", "-"):
        time_part = time_part.split(separator, 1)[0]
    return time_part[:8]


def _write_sheet(workbook: Workbook, title: str, rows: Sequence[Mapping[str, object]], apply_values: Mapping[str, object] | None = None) -> None:
    if title in workbook.sheetnames:
        del workbook[title]
    ws = workbook.create_sheet(title)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B22222")
    url_column = HEADERS.index("Tesla Job URL") + 1
    for row in rows:
        ws.append([_date_only(row["first_seen_at"]), _time_only(row["first_seen_at"]), row["posted_at"], row["title"],
                   row["location"], row["category"], row["employment_type"], row["job_id"], row["url"],
                   (apply_values or {}).get(str(row["job_id"]), "")])
        link_cell = ws.cell(ws.max_row, url_column)
        link_cell.hyperlink = str(row["url"])
        link_cell.style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in {"A": 23, "B": 15, "C": 23, "D": 48, "E": 32, "F": 28, "G": 18, "H": 14, "I": 70, "J": 14}.items():
        ws.column_dimensions[column].width = width


def export_workbook(path: Path, new_rows: Sequence[Mapping[str, object]], all_rows: Sequence[Mapping[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(path) if path.exists() else Workbook()
    if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) == 1:
        del workbook["Sheet"]
    existing_apply: dict[str, object] = {}
    if "New Jobs" in workbook.sheetnames:
        old_sheet = workbook["New Jobs"]
        # Locate columns by header name so preserved "Apply?" values survive
        # layout changes (older sheets lack the "Time Discovered" column).
        header = [cell.value for cell in old_sheet[1]]
        try:
            id_index, apply_index = header.index("Req ID"), header.index("Apply?")
        except ValueError:
            id_index = apply_index = None
        if id_index is not None and apply_index is not None:
            for row in old_sheet.iter_rows(min_row=2, values_only=True):
                if len(row) > max(id_index, apply_index) and row[id_index] is not None:
                    existing_apply[str(row[id_index])] = row[apply_index]
    _write_sheet(workbook, "New Jobs", new_rows, existing_apply)
    _write_sheet(workbook, "All Jobs", all_rows)
    workbook.save(path)
