from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.database import JobDatabase
from src.excel_exporter import export_workbook
from src.parser import TeslaSchemaError, parse_jobs
from src.tesla_client import TeslaClient, TeslaFetchError


PAYLOAD = {
    "lookup": {"locations": {"1": "Palo Alto, California", "2": "Berlin, Germany"},
               "departments": {"1": "Tesla AI"}, "types": {"1": "Full-Time", "2": "Intern/Apprentice"}},
    "listings": [{"id": 100, "t": "Engineer", "l": 1, "dp": 1, "y": 1},
                 {"id": 101, "t": "German Engineer", "l": 2, "dp": 1, "y": 1},
                 {"id": 102, "t": "Intern", "l": 1, "dp": 1, "y": 2}],
}


def test_parser_filters_us_full_time_and_missing_optional_fields():
    jobs = parse_jobs(PAYLOAD)
    assert [job.job_id for job in jobs] == ["100"]
    assert jobs[0].posted_at is None
    assert jobs[0].url.endswith("tjt-100")


def test_parser_empty_response_and_duplicates():
    assert parse_jobs({"lookup": {}, "listings": []}) == []
    duplicate = {**PAYLOAD, "listings": [PAYLOAD["listings"][0], PAYLOAD["listings"][0]]}
    with pytest.raises(TeslaSchemaError, match="duplicate"):
        parse_jobs(duplicate)


def test_database_new_detection_duplicate_and_updates(tmp_path: Path):
    db = JobDatabase(tmp_path / "jobs.db"); db.initialize()
    jobs = parse_jobs(PAYLOAD)
    db.replace_baseline(jobs)
    new, existing = db.record_jobs(jobs)
    assert new == [] and existing == 1
    payload2 = {**PAYLOAD, "listings": PAYLOAD["listings"] + [{"id": 103, "t": "New Role", "l": 1, "dp": 1, "y": 1}]}
    new, existing = db.record_jobs(parse_jobs(payload2))
    assert [row["job_id"] for row in new] == ["103"] and existing == 1
    assert [row["job_id"] for row in db.new_job_history()] == ["103"]
    with db.connect() as con:
        assert con.execute("SELECT is_active FROM jobs WHERE job_id='100'").fetchone()[0] == 1


def test_baseline_discovery_date_can_be_set_without_touching_new_jobs(tmp_path: Path):
    db = JobDatabase(tmp_path / "jobs.db"); db.initialize()
    db.replace_baseline(parse_jobs(PAYLOAD), "2026-08-01")
    payload2 = {**PAYLOAD, "listings": PAYLOAD["listings"] + [{"id": 103, "t": "New Role", "l": 1, "dp": 1, "y": 1}]}
    db.record_jobs(parse_jobs(payload2))
    assert db.set_baseline_discovery_date("2026-07-31") == 1
    with db.connect() as con:
        assert con.execute("SELECT first_seen_at FROM jobs WHERE job_id='100'").fetchone()[0] == "2026-07-31"
        assert con.execute("SELECT first_seen_at FROM jobs WHERE job_id='103'").fetchone()[0] != "2026-07-31"


def test_excel_generation_has_hyperlink(tmp_path: Path):
    db = JobDatabase(tmp_path / "jobs.db"); db.initialize()
    db.replace_baseline(parse_jobs(PAYLOAD))
    path = tmp_path / "Tesla_New_Jobs.xlsx"
    export_workbook(path, [], db.current_jobs())
    book = load_workbook(path)
    assert book.sheetnames == ["New Jobs", "All Jobs"]
    assert book["All Jobs"]["A1"].value == "Date Discovered"
    assert book["All Jobs"]["A2"].value == str(db.current_jobs()[0]["first_seen_at"])[:10]
    assert book["All Jobs"]["H2"].hyperlink.target.endswith("tjt-100")


def test_client_handles_http_error():
    class Response:
        status_code = 500
        text = "broken"
        def raise_for_status(self): raise __import__('requests').HTTPError("500")
    class Session:
        def get(self, *args, **kwargs): return Response()
    with pytest.raises(TeslaFetchError, match="Unable to fetch"):
        TeslaClient(session=Session()).fetch_jobs()
